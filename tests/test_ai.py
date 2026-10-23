import asyncio, json, sys, random
from playwright.async_api import async_playwright
import openpyxl

def mk():
    single=[]
    for i in range(10):
        single.append({"stem":f"单选测试题{i+1}，下列正确的是（  ）","options":[f"错误甲{i}",f"正确项{i}",f"错误乙{i}",f"错误丙{i}"],"answer":"B","verify":(i%3==0)})
    multi=[]
    pat=["AB","ABC","ACD","ABCD","BC"]
    for i in range(5):
        a=pat[i]; opts=[f"选项{c}{i}" for c in "ABCD"]
        multi.append({"stem":f"多选测试题{i+1}（  ）","options":opts,"answer":a,"verify":False})
    tf=[{"stem":f"判断测试{i+1}。（  ）","answer":"正确" if i%2==0 else "错误","verify":False} for i in range(5)]
    return {"single":single,"multi":multi,"tf":tf}

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome', args=['--no-sandbox'])
        ctx = await b.new_context(accept_downloads=True, viewport={'width':1100,'height':900})
        pg = await ctx.new_page()
        errs=[]
        pg.on('pageerror', lambda e: errs.append(str(e)))
        pg.on('console', lambda m: errs.append(m.text) if m.type=='error' else None)
        mode={'v':'ok'}
        async def handler(route):
            req = route.request
            hdr = {'Access-Control-Allow-Origin':'*','Access-Control-Allow-Headers':'*','Access-Control-Allow-Methods':'*'}
            if req.method=='OPTIONS':
                await route.fulfill(status=204, headers=hdr); return
            if mode['v']=='401':
                await route.fulfill(status=401, headers=hdr, body='{"error":"bad key"}'); return
            if req.url.endswith('/models'):
                await route.fulfill(status=200, headers=hdr, body='{"data":[]}'); return
            body = json.loads(req.post_data)
            assert body['response_format']['type']=='json_object', body
            assert 'Bearer sk-test' == req.headers.get('authorization'), req.headers
            prompt = body['messages'][1]['content']
            assert '胸腔闭式引流' in prompt and 'JSON' in prompt
            content = '```json\n'+json.dumps(mk(), ensure_ascii=False)+'\n```'
            await route.fulfill(status=200, headers={**hdr,'Content-Type':'application/json'}, body=json.dumps({"choices":[{"message":{"content":content}}]}))
        await pg.route('https://api.deepseek.com/**', handler)
        await pg.goto('file:///home/claude/work/index.html')

        # 1. no key -> error message
        await pg.fill('#aiTopic','胸腔闭式引流护理团体标准基础知识')
        await pg.click('#aiGenBtn')
        print('NOKEY:', (await pg.inner_text('#aiStatus'))[:60])

        # 401
        await pg.fill('#aiKey','sk-bad'); mode['v']='401'
        await pg.click('#aiGenBtn'); await pg.wait_for_function("document.getElementById('aiStatus').textContent.includes('✕')")
        print('401:', (await pg.inner_text('#aiStatus'))[:60])
        mode['v']='ok'
        # test conn
        await pg.fill('#aiKey','sk-test')
        await pg.click('#aiTestBtn'); await pg.wait_for_function("document.getElementById('aiStatus').textContent.includes('连接成功')")
        print('TEST ok')

        await pg.click('#aiGenBtn')
        await pg.wait_for_function("document.getElementById('aiStatus').textContent.includes('已生成')", timeout=20000)
        print('GEN:', (await pg.inner_text('#aiStatus')).replace('\n',' | '))
        # answer distribution for single
        ans = await pg.evaluate("""() => {
          const blocks=[...document.querySelectorAll('.qtype-block')];
          const sb=blocks[0]; const out=[];
          sb.querySelectorAll('.qrow.qedit').forEach(r=>{
            const rows=[...r.querySelectorAll('.opt-row')];
            rows.forEach((o,i)=>{ if(o.querySelector('input').checked) out.push('ABCD'[i]); });
          }); return out; }""")
        print('SINGLE ANSWERS', ans, {c:ans.count(c) for c in 'ABCD'})
        mans = await pg.evaluate("""() => { const sb=document.querySelectorAll('.qtype-block')[1]; const out=[];
          sb.querySelectorAll('.qrow.qedit').forEach(r=>{ let s=''; r.querySelectorAll('.opt-row').forEach((o,i)=>{ if(o.querySelector('input').checked) s+='ABCDE'[i]; }); out.push(s); }); return out; }""")
        print('MULTI ANSWERS', mans)
        print('SUMMARY:', (await pg.inner_text('#issueSummary')).replace('\n',' | ')[:300])

        # edit: delete a single question -> expect issue about count
        await pg.click('.qtype-block >> nth=0 >> .qdel >> nth=0')
        print('AFTER DELETE:', (await pg.inner_text('#issueSummary')).replace('\n',' | ')[:200])
        # add blank -> flagged
        await pg.click('.qtype-block >> nth=0 >> text=添加一题')
        print('AFTER ADD BLANK:', (await pg.inner_text('#issueSummary')).replace('\n',' | ')[:260])
        # fill the blank question
        blk = pg.locator('.qtype-block').nth(0).locator('.qrow.qedit').last
        await blk.locator('.qstem-edit').fill('手动添加的单选题（  ）')
        texts = blk.locator('.opt-text')
        for i,t in enumerate(['甲','乙','丙','丁']): await texts.nth(i).fill(t)
        await blk.locator('.opt-row input').nth(2).check()
        print('AFTER FILL:', (await pg.inner_text('#issueSummary')).replace('\n',' | ')[:200])

        async with pg.expect_download() as dl:
            await pg.click('#exportBtn')
        d = await dl.value
        path='/home/claude/work/tests/out_ai.xlsx'; await d.save_as(path)
        print('DOWNLOAD', d.suggested_filename)
        wb=openpyxl.load_workbook(path)
        print(wb.sheetnames)
        for n in ['试卷信息','单选题','多选题','是非题']:
            ws=wb[n]; print(n, ws.max_row-1, [c.value for c in ws[2]][:9])
        print('ERRORS', errs)
        await pg.screenshot(path='/home/claude/work/tests/shot_ai.png', full_page=True)
        await b.close()
asyncio.run(main())
