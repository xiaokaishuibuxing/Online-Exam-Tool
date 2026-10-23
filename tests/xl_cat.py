import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1])
ws=wb['多选题']; print('multi answers', [r[9].value for r in ws.iter_rows(min_row=2)])
ws=wb['单选题']; print('single answers', [r[7].value for r in ws.iter_rows(min_row=2)])
ws=wb['是非题']; print('tf answers', [r[1].value for r in ws.iter_rows(min_row=2)])
